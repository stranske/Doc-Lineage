"""Word and PowerPoint extraction returning the same span shape."""

from __future__ import annotations

import hashlib
from typing import TYPE_CHECKING

from doc_lineage.extract.models import CoverageStats, Document, Span

if TYPE_CHECKING:
    from pathlib import Path

    from doc_lineage.extract.cache import ExtractCache

#: Office documents are cached as one whole-document entry; page ``0`` is the
#: sentinel pointer for "not a page of this document, but the document itself".
_WHOLE_DOCUMENT_PAGE = 0


def extract_docx(path: Path, stable_id: str, *, cache: ExtractCache | None = None) -> Document:
    spans = _cached_spans(cache, stable_id)
    if spans is None:
        from docx import Document as DocxDocument

        doc = DocxDocument(str(path))
        spans = []
        for index, paragraph in enumerate(doc.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            spans.append(
                Span(
                    text=text,
                    page=index,
                    bbox=None,
                    source="text_layer",
                )
            )
        _store_spans(cache, stable_id, spans)
    return _document(stable_id, spans)


def extract_pptx(path: Path, stable_id: str, *, cache: ExtractCache | None = None) -> Document:
    spans = _cached_spans(cache, stable_id)
    if spans is None:
        from pptx import Presentation

        presentation = Presentation(str(path))
        spans = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue
                text = shape.text.strip()
                if text:
                    texts.append(text)
            if texts:
                spans.append(
                    Span(
                        text="\n".join(texts),
                        page=slide_index,
                        bbox=None,
                        source="text_layer",
                    )
                )
        _store_spans(cache, stable_id, spans)
    return _document(stable_id, spans)


def _cached_spans(cache: ExtractCache | None, stable_id: str) -> list[Span] | None:
    if cache is None:
        return None
    return cache.get(stable_id, _WHOLE_DOCUMENT_PAGE, "office")


def _store_spans(cache: ExtractCache | None, stable_id: str, spans: list[Span]) -> None:
    if cache is None:
        return
    cache.put(stable_id, _WHOLE_DOCUMENT_PAGE, spans, "office")


def _document(stable_id: str, spans: list[Span]) -> Document:
    coverage = CoverageStats(
        pages_with_text_layer=len(spans),
        pages_recognized=0,
        pages_unreadable=0,
    )
    return Document(stable_id=stable_id, spans=spans, coverage=coverage)


def stable_id_for_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()
