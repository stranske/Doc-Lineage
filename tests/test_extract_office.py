"""Word and PowerPoint extraction tests."""

from __future__ import annotations

import hashlib
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from doc_lineage.extract import ExtractCache, extract, reset_cache


@pytest.fixture(autouse=True)
def _isolated_module_cache(tmp_path: Path) -> None:
    reset_cache(ExtractCache(tmp_path / "module-cache"))
    yield
    reset_cache(None)


def _write_docx(path: Path, paragraphs: list[str]) -> None:
    doc = DocxDocument()
    for text in paragraphs:
        doc.add_paragraph(text)
    doc.save(str(path))


def test_extract_docx_returns_paragraph_spans_and_pointers(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _write_docx(path, ["First paragraph", "", "Second paragraph"])

    document = extract(path, cache=ExtractCache(tmp_path / "cache"))

    assert [(span.page, span.text, span.source, span.bbox) for span in document.spans] == [
        (1, "First paragraph", "text_layer", None),
        (3, "Second paragraph", "text_layer", None),
    ]
    assert document.coverage.pages_with_text_layer == 2
    assert document.coverage.pages_recognized == 0
    assert document.coverage.pages_unreadable == 0


def test_extract_pptx_returns_slide_spans_and_pointers(tmp_path: Path) -> None:
    path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Slide title"
    presentation.save(str(path))

    document = extract(path, cache=ExtractCache(tmp_path / "cache2"))

    assert [(span.page, span.text, span.source, span.bbox) for span in document.spans] == [
        (1, "Slide title", "text_layer", None)
    ]
    assert document.coverage.pages_with_text_layer == 1


def test_office_extraction_uses_the_cache(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """A second read of the same content must not re-parse the file."""
    path = tmp_path / "cached.docx"
    _write_docx(path, ["Cached paragraph"])
    active_cache = ExtractCache(tmp_path / "cache3")

    first = extract(path, cache=active_cache)

    def _fail(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("docx was re-parsed instead of served from the cache")

    monkeypatch.setattr("docx.Document", _fail)

    second = extract(path, cache=active_cache)

    assert [span.text for span in second.spans] == [span.text for span in first.spans]
    assert second.stable_id == first.stable_id
    assert second.coverage == first.coverage


def test_legacy_binary_office_formats_are_rejected(tmp_path: Path) -> None:
    """``python-pptx``/``python-docx`` read OOXML only; say so instead of failing opaquely."""
    legacy_ppt = tmp_path / "deck.ppt"
    legacy_ppt.write_bytes(b"\xd0\xcf\x11\xe0legacy-ole-container")
    legacy_doc = tmp_path / "memo.doc"
    legacy_doc.write_bytes(b"\xd0\xcf\x11\xe0legacy-ole-container")

    with pytest.raises(ValueError, match=r"\.pptx"):
        extract(legacy_ppt, cache=ExtractCache(tmp_path / "cache4"))
    with pytest.raises(ValueError, match=r"\.docx"):
        extract(legacy_doc, cache=ExtractCache(tmp_path / "cache5"))


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
@pytest.mark.parametrize("reopen_cache", [False, True], ids=["shared-memory", "reopened-disk"])
def test_office_replacement_keeps_snapshot_identity_and_cached_text(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, suffix: str, reopen_cache: bool
) -> None:
    path = tmp_path / f"snapshot{suffix}"

    def write_document(text: str) -> None:
        if suffix == ".docx":
            _write_docx(path, [text])
        else:
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = text
            presentation.save(str(path))

    write_document("Original snapshot")
    original = path.read_bytes()
    write_document("Revised on disk")
    revised = path.read_bytes()
    path.write_bytes(original)
    read_bytes = Path.read_bytes

    def replace_after_read(candidate: Path) -> bytes:
        content = read_bytes(candidate)
        if candidate == path:
            candidate.write_bytes(revised)
        return content

    cache_root = tmp_path / "snapshot-cache"
    active_cache = ExtractCache(cache_root if reopen_cache else None)
    with monkeypatch.context() as patch:
        patch.setattr(Path, "read_bytes", replace_after_read)
        first = extract(path, cache=active_cache)

    assert path.read_bytes() == revised
    assert first.stable_id == hashlib.sha256(original).hexdigest()
    assert [span.text for span in first.spans] == ["Original snapshot"]
    path.write_bytes(original)

    def fail_reparse(*_args: object, **_kwargs: object) -> None:
        raise AssertionError("original snapshot must be served from its cache")

    monkeypatch.setattr("docx.Document" if suffix == ".docx" else "pptx.Presentation", fail_reparse)
    replay_cache = ExtractCache(cache_root) if reopen_cache else active_cache
    replay = extract(path, cache=replay_cache)
    assert replay.stable_id == first.stable_id
    assert replay.spans == first.spans
    assert [span.text for span in replay.spans] == ["Original snapshot"]
